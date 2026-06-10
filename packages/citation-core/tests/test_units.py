"""Birlik testlari: saralash/collation, translit, imlo, DOCX, preprocessing."""

from __future__ import annotations

import io

from citation_core import Pipeline, collation_key, sort_sources, split_sources
from citation_core.exporters import export_docx, export_text
from citation_core.sorting import group_of
from citation_core.spell import SpellChecker
from citation_core.translit import Transliterator

pipe = Pipeline()


# ---------- Collation: kirill-lotin aralash alifbo ----------

def test_collation_mixed_alphabet() -> None:
    # Кирилл «Абдуллаев» va lotin «Aliyev» bir guruhda, A harfida turadi
    names = ["G'ulomov", "Абдуллаев", "Aliyev", "Saidov", "Бобоев", "Karimov"]
    ordered = sorted(names, key=collation_key)
    assert ordered.index("Абдуллаев") < ordered.index("Aliyev") < ordered.index("Бобоев")
    assert ordered.index("Бобоев") < ordered.index("Karimov") < ordered.index("Saidov")
    assert ordered[-1] == "G'ulomov"  # o', g' — alifbo oxirida (z dan keyin)


def test_oak_three_groups() -> None:
    texts = [
        "Jahon banki ma'lumotlari [Elektron resurs]. – URL: https://data.worldbank.org (murojaat sanasi: 15.05.2026).",
        "Karimov A.A. Iqtisodiyot nazariyasi: Darslik. – Toshkent: Iqtisod-moliya, 2020. – 350 b.",
        "O'zbekiston Respublikasining «Ta'lim to'g'risida»gi Qonuni. O'RQ-637-son, 23.09.2020 // Qonunchilik ma'lumotlari milliy bazasi, lex.uz",
        "Smith J., Brown K. Machine Learning in Economics // Journal of Economic Perspectives. – 2023. – Vol. 37, №2. – P. 115–142.",
        "Абдуллаев Т.Т. Молия назарияси: Дарслик. – Тошкент: Иқтисод-молия, 2019. – 312 б.",
    ]
    sources = [pipe.parse_one(t) for t in texts]
    ordered = sort_sources(sources)
    groups = [group_of(s) for s in ordered]
    assert groups == sorted(groups), "Guruhlar tartibi: 1 → 2 → 3"
    # 2-guruh ichida: kirill/lotin o'zbek-rus avval, xorijiy keyin
    main = [s for s in ordered if group_of(s) == 2]
    surnames = [s.fields.authors[0].surname for s in main]
    assert surnames.index("Абдуллаев") < surnames.index("Karimov") < surnames.index("Smith")


# ---------- Transliteratsiya (2021 qoidalari) ----------

def test_translit_cyr_to_lat() -> None:
    tr = Transliterator()
    assert tr.cyr_to_lat("Тошкент") == "Toshkent"
    assert tr.cyr_to_lat("Ўзбекистон") == "O'zbekiston"
    assert tr.cyr_to_lat("ҳақида") == "haqida"
    # е: so'z boshida «ye», undoshdan keyin «e»
    assert tr.cyr_to_lat("Европа") == "Yevropa"
    assert tr.cyr_to_lat("мактаб") == "maktab"


def test_translit_lat_to_cyr() -> None:
    tr = Transliterator()
    assert tr.lat_to_cyr("Toshkent") == "Тошкент"
    assert tr.lat_to_cyr("o'zbek") == "ўзбек"
    assert tr.lat_to_cyr("shahar") == "шаҳар"


def test_translit_protects_urls_and_exceptions() -> None:
    tr = Transliterator(exceptions={"Google"})
    out = tr.cyr_to_lat("Манба: https://лекс.uz/докс ва Google сайти")
    assert "https://лекс.uz/докс" in out  # URL tegilmaydi
    assert "Google" in out


# ---------- Imlo ----------

def test_spell_common_fix_and_apply() -> None:
    sc = SpellChecker()
    issues = sc.check("raqamli tehnologiya rivoji")
    assert any(i.xato == "tehnologiya" and "texnologiya" in i.taklif for i in issues)
    fixed = SpellChecker.apply("raqamli tehnologiya rivoji", issues)
    assert fixed == "raqamli texnologiya rivoji"


def test_spell_custom_dictionary_whitelist() -> None:
    sc = SpellChecker(custom_words={"tehnologiya"})  # atama sifatida qo'shilgan
    assert sc.check("tehnologiya") == []


# ---------- DOCX eksport ----------

def test_docx_export_structure() -> None:
    from docx import Document

    texts = [
        "O'zbekiston Respublikasining «Ta'lim to'g'risida»gi Qonuni. O'RQ-637-son, 23.09.2020 // Qonunchilik ma'lumotlari milliy bazasi, lex.uz",
        "Karimov A.A. Iqtisodiyot nazariyasi: Darslik. – Toshkent: Iqtisod-moliya, 2020. – 350 b.",
        "Jahon banki ma'lumotlari [Elektron resurs]. – URL: https://data.worldbank.org (murojaat sanasi: 15.05.2026).",
    ]
    sources = [pipe.parse_one(t) for t in texts]
    data = export_docx(sources)
    doc = Document(io.BytesIO(data))
    paras = [p.text for p in doc.paragraphs if p.text.strip()]
    assert paras[0] == "FOYDALANILGAN ADABIYOTLAR RO'YXATI"
    assert any(p.startswith("I. Normativ-huquqiy") for p in paras)
    assert any(p.startswith("II. Asosiy adabiyotlar") for p in paras)
    assert any(p.startswith("III. Internet manbalar") for p in paras)
    # Raqamlash uzluksiz: 1., 2., 3.
    numbered = [p for p in paras if p[0].isdigit()]
    assert [p.split(".")[0] for p in numbered] == ["1", "2", "3"]
    # Shrift: Times New Roman 14
    run = next(p for p in doc.paragraphs if p.text.startswith("1.")).runs[0]
    assert run.font.name == "Times New Roman"
    assert run.font.size.pt == 14
    # Osma xat (hanging indent)
    pf = next(p for p in doc.paragraphs if p.text.startswith("1.")).paragraph_format
    assert round(pf.left_indent.inches, 2) == 0.5
    assert round(pf.first_line_indent.inches, 2) == -0.5


def test_export_text_clipboard() -> None:
    texts = [
        "Karimov A.A. Iqtisodiyot nazariyasi: Darslik. – Toshkent: Iqtisod-moliya, 2020. – 350 b.",
    ]
    out = export_text([pipe.parse_one(t) for t in texts])
    assert "1. Karimov A.A." in out


# ---------- Preprocessing ----------

def test_split_sources_multiline_and_wrapped() -> None:
    text = (
        "1. Karimov A.A. Iqtisodiyot nazariyasi: Darslik. – Toshkent: Iqtisod-moliya, 2020. – 350 b.\n"
        "2. Saidov B.B. Raqamli iqtisodiyotda innovatsiyalar // Iqtisodiyot va\n"
        "ta'lim. – Toshkent, 2021. – №3. – B. 45–52.\n"
        "2. Saidov B.B. Raqamli iqtisodiyotda innovatsiyalar // Iqtisodiyot va ta'lim. – Toshkent, 2021. – №3. – B. 45–52.\n"
    )
    parts = split_sources(text)
    assert len(parts) == 2  # o'ralgan satr qo'shildi, takror olib tashlandi
    assert parts[1].endswith("B. 45–52.")


def test_pipeline_warnings_on_bad_year() -> None:
    src = pipe.parse_one("Karimov A.A. Test kitob. – Toshkent: Fan, 1200. – 100 b.")
    # 1200 yil YEAR regexga tushmaydi — yil topilmaydi yoki ogohlantiriladi
    assert src.fields.year is None or src.warnings
